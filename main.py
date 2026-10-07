#!/usr/bin/env python3
"""
Extract calculated fields from a Tableau workbook XML and export them to Excel.

Usage:
    python3 main.py --source-xml tableau-def.xml
    python3 main.py --source-xml "My Workbook.xml" --output custom_name.xlsx
"""
import argparse
import html
import re
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.styles.borders import Side
from openpyxl.utils import get_column_letter

COLUMN_WIDTHS = {
    'Calculated Field Name': 30,
    'Internal Name': 40,
    'Datasource': 25,
    'Data Type': 12,
    'Calculation Logic (Tableau)': 60,
    'Dependencies': 45,
    'Used on Chart': 14,
    'Worksheet(s)': 40,
    'Dashboard Page(s)': 35,
}


# ===== CLI =====
def parse_args():
    parser = argparse.ArgumentParser(
        description='Export Tableau calculated fields from a workbook XML to an Excel file.'
    )
    parser.add_argument('--source-xml', required=True,
                        help='Path to the XML extracted from Tableau.')
    parser.add_argument('--output', default=None,
                        help='Output .xlsx path. Defaults to the source file name, '
                             'lowercased with spaces replaced by underscores.')
    return parser.parse_args()


def default_output_path(source_xml: Path) -> Path:
    """'My Workbook Calc Fields.xml' -> 'my_workbook_calc_fields.xlsx' (same folder)."""
    name = source_xml.stem.lower().replace(' ', '_')
    return source_xml.with_name(f'{name}.xlsx')


# ===== Parsing =====
def clean_formula(formula: str) -> str:
    return html.unescape(formula).replace('\r\n', '\n').replace('\r', '\n')


def build_calc_fields(root):
    calc_fields = {}
    for ds in root.findall('.//datasource'):
        ds_name = ds.get('name', '')
        ds_caption = ds.get('caption', ds_name)

        for calc in ds.findall('.//calculations/calculation'):
            col_name = calc.get('column', '')
            formula = calc.get('formula', '')
            if formula:
                calc_fields[col_name] = {
                    'internal_name': col_name, 'caption': '', 'formula': clean_formula(formula),
                    'datasource': ds_caption, 'data_type': '',
                }

        for col in ds.findall('.//column'):
            col_name = col.get('name', '')
            col_caption = col.get('caption', '')
            calc_elem = col.find('calculation')
            if calc_elem is not None:
                formula = calc_elem.get('formula', '')
                if formula and col_name not in calc_fields:
                    calc_fields[col_name] = {
                        'internal_name': col_name, 'caption': col_caption,
                        'formula': clean_formula(formula), 'datasource': ds_caption,
                        'data_type': col.get('datatype', ''),
                    }
            if col_name in calc_fields and col_caption:
                calc_fields[col_name]['caption'] = col_caption

        for mr in ds.findall('.//metadata-record'):
            ln, cap, lt = mr.find('local-name'), mr.find('caption'), mr.find('local-type')
            if ln is not None and ln.text and ln.text in calc_fields:
                if cap is not None and cap.text:
                    calc_fields[ln.text]['caption'] = cap.text
                if lt is not None and lt.text:
                    calc_fields[ln.text]['data_type'] = lt.text
    return calc_fields


def extract_dependencies(formula):
    refs = re.findall(r'\[([^\]]+)\]', formula)
    return sorted({ref for ref in refs if ref != 'Parameters'})


def build_worksheet_fields(root):
    worksheet_fields = {}
    for ws in root.findall('.//worksheet'):
        fields_used = set()
        for dep in ws.findall('.//datasource-dependencies'):
            for col in dep.findall('column'):
                if col.get('name'):
                    fields_used.add(col.get('name'))
        for ci in ws.findall('.//column-instance'):
            if ci.get('column'):
                fields_used.add(ci.get('column'))
        worksheet_fields[ws.get('name', '')] = fields_used
    return worksheet_fields


def build_dashboard_worksheets(root):
    ws_names = {ws.get('name') for ws in root.findall('.//worksheet')}
    dashboard_worksheets = {}
    for dashboard in root.findall('.//dashboard'):
        dashboard_worksheets[dashboard.get('name', '')] = {
            zone.get('name', '') for zone in dashboard.iter('zone')
            if zone.get('name', '') in ws_names
        }
    return dashboard_worksheets


def build_dataframe(calc_fields, worksheet_fields, dashboard_worksheets):
    name_to_caption = {k.strip('[]'): (v['caption'] or k.strip('[]')) for k, v in calc_fields.items()}

    ws_to_dashboard = {}
    for dn, wss in dashboard_worksheets.items():
        for w in wss:
            ws_to_dashboard.setdefault(w, []).append(dn)

    rows = []
    for field_key, info in calc_fields.items():
        deps = [name_to_caption.get(d, d) for d in extract_dependencies(info['formula'])]
        used_ws = sorted(ws for ws, flds in worksheet_fields.items() if field_key in flds)
        used_dash = sorted({d for ws in used_ws for d in ws_to_dashboard.get(ws, [])})
        rows.append({
            'Calculated Field Name': info['caption'] or field_key.strip('[]'),
            'Internal Name': field_key.strip('[]'),
            'Datasource': info['datasource'],
            'Data Type': info['data_type'],
            'Calculation Logic (Tableau)': info['formula'],
            'Dependencies': ', '.join(deps) if deps else 'None (static value)',
            'Used on Chart': 'Y' if used_ws else 'N',
            'Worksheet(s)': '\n'.join(used_ws),
            'Dashboard Page(s)': '\n'.join(used_dash),
        })

    df = pd.DataFrame(rows, columns=list(COLUMN_WIDTHS.keys()))
    if not df.empty:
        df = df.sort_values(by=['Used on Chart', 'Datasource', 'Calculated Field Name'],
                            ascending=[False, True, True])
    return df


# ===== Excel output =====
def write_excel(df, output_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Calculated Fields'

    headers = list(df.columns)
    header_font = Font(bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(start_color='2F5496', end_color='2F5496', fill_type='solid')
    thin = Side(style='thin')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(vertical='top', wrap_text=True)
    green = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
    red = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border

    for row_idx, (_, row) in enumerate(df.iterrows(), 2):
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=row[header])
            cell.alignment = wrap
            cell.border = border
            if header == 'Used on Chart':
                cell.fill = green if row[header] == 'Y' else red

    for col_idx, header in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = COLUMN_WIDTHS.get(header, 20)

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{len(df) + 1}'
    wb.save(output_path)


def print_summary(df, output_path, worksheet_fields, dashboard_worksheets):
    print('Excel file saved successfully!')
    print(f'Path: {output_path}')
    print('\nSummary:')
    print(f'  Total calculated fields: {len(df)}')
    print(f"  Fields used on charts: {(df['Used on Chart'] == 'Y').sum()}")
    print(f"  Fields NOT used on charts: {(df['Used on Chart'] == 'N').sum()}")
    print(f'\n  Dashboard pages: {len(dashboard_worksheets)}')
    for dn in sorted(dashboard_worksheets):
        print(f'    - {dn} ({len(dashboard_worksheets[dn])} charts)')
    print(f'\n  Total worksheets (charts): {len(worksheet_fields)}')


def main():
    args = parse_args()
    source_xml = Path(args.source_xml)
    if not source_xml.is_file():
        sys.exit(f'Error: source XML not found: {source_xml}')

    output_path = Path(args.output) if args.output else default_output_path(source_xml)

    try:
        root = ET.parse(source_xml).getroot()
    except ET.ParseError as e:
        sys.exit(f'Error: could not parse XML ({e})')

    calc_fields = build_calc_fields(root)
    worksheet_fields = build_worksheet_fields(root)
    dashboard_worksheets = build_dashboard_worksheets(root)
    df = build_dataframe(calc_fields, worksheet_fields, dashboard_worksheets)

    write_excel(df, output_path)
    print_summary(df, output_path, worksheet_fields, dashboard_worksheets)


if __name__ == '__main__':
    main()